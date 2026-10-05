import tempfile
import unittest
from contextlib import closing
from pathlib import Path
from unittest.mock import patch

from fastapi.testclient import TestClient

import main
import office
import tables

LOAN_EXPORT = (
    "Loan Report\nRun date: 2026-09-01\n\n"
    "Account No,Branch,Status,Balance,Opened\n"
    "0012345,Lae,Open,\"1,200.50\",01/02/2024\n"
    "0012346,Lae,Closed,300,15/03/2024\n"
    "0012347,Madang,Open,,20/12/2023\n"
    "0012348,Madang,Open,500,05/01/2024\n"
)


class SpreadsheetEngineTests(unittest.TestCase):
    def setUp(self) -> None:
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)

    def tearDown(self) -> None:
        self.temp.cleanup()

    def csv(self, text: str = LOAN_EXPORT, name: str = "loans.csv", encoding: str = "utf-8") -> dict:
        path = self.root / name
        path.write_bytes(text.encode(encoding))
        return tables.load_sheets(path)

    def test_title_rows_are_skipped_and_columns_are_typed(self) -> None:
        frame = self.csv()["Sheet1"]
        self.assertEqual(list(frame.columns), ["Account No", "Branch", "Status", "Balance", "Opened"])
        self.assertEqual(len(frame), 4)
        self.assertEqual(tables.column_type(frame["Balance"]), "number")
        self.assertEqual(tables.column_type(frame["Opened"]), "date")
        # Identifiers keep their leading zeros instead of becoming numbers.
        self.assertEqual(tables.column_type(frame["Account No"]), "text")
        self.assertEqual(frame["Account No"].iloc[0], "0012345")
        # 15/03/2024 shows the export is day-first, so 01/02/2024 is 1 February.
        self.assertEqual(frame["Opened"].iloc[0].date().isoformat(), "2024-02-01")

    def test_letter_prefixed_identifiers_stay_text_and_currency_amounts_parse(self) -> None:
        frame = self.csv(
            "Case ID,Amount\nC0042,\"₹1,20,000\"\nC0043,INR 500\nC0044,$12.50\nC0045,(300)\n"
        )["Sheet1"]
        self.assertEqual(tables.column_type(frame["Case ID"]), "text")
        self.assertEqual(frame["Case ID"].tolist(), ["C0042", "C0043", "C0044", "C0045"])
        self.assertEqual(frame["Amount"].tolist(), [120000.0, 500.0, 12.5, -300.0])

    def test_utf16_tab_delimited_excel_export_is_read(self) -> None:
        text = "Branch\tCases\r\nLae\t4\r\nMadang\t7\r\n"
        frame = self.csv("﻿" + text, "export.csv", "utf-16-le")["Sheet1"]
        self.assertEqual(list(frame.columns), ["Branch", "Cases"])
        self.assertEqual(frame["Cases"].sum(), 11)

    def test_every_excel_sheet_is_loaded(self) -> None:
        from openpyxl import Workbook

        workbook = Workbook()
        first = workbook.active
        first.title = "KYC"
        first.append(["Customer", "Status"])
        first.append(["Ravi", "Pending"])
        second = workbook.create_sheet("Dormant")
        second.append(["Quarterly dormant list"])
        second.append(["Account", "Balance"])
        second.append(["A1", 10])
        second.append(["A2", 25.5])
        path = self.root / "tracker.xlsx"
        workbook.save(path)
        sheets = tables.load_sheets(path)
        self.assertEqual(list(sheets), ["KYC", "Dormant"])
        self.assertEqual(list(sheets["Dormant"].columns), ["Account", "Balance"])
        result = tables.execute(sheets, "Dormant", {"metrics": [{"function": "sum", "column": "Balance"}]})
        self.assertEqual(result["rows"], [["35.50"]])

    def test_grouped_count_and_sum_report_excluded_empty_values(self) -> None:
        sheets = self.csv()
        result = tables.execute(sheets, "Sheet1", {
            "filters": [{"column": "Status", "operator": "equals", "value": "OPEN"}],
            "groupBy": ["Branch"],
            "metrics": [{"function": "count"}, {"function": "sum", "column": "Balance"}],
        })
        self.assertEqual(result["columns"], ["Branch", "Count", "Sum of Balance"])
        self.assertEqual(result["rows"], [["Madang", "2", "500"], ["Lae", "1", "1,200.50"]])
        self.assertEqual(result["matchedRows"], 3)
        self.assertEqual(result["excludedEmpty"], {"Balance": 1})
        text = tables.describe(result, "loans.csv", "Sheet1", False)
        self.assertIn('Filters: Status is "OPEN".', text)
        self.assertIn("Computed over 3 of 4 rows in loans.csv.", text)
        self.assertIn("1 matching row with an empty Balance was left out", text)

    def test_record_lookup_date_and_numeric_filters(self) -> None:
        sheets = self.csv()
        lookup = tables.execute(sheets, "Sheet1", {
            "filters": [{"column": "Account No", "operator": "equals", "value": "0012346"}],
            "columns": ["Branch", "Balance"],
        })
        self.assertEqual(lookup["rows"], [["Lae", "300"]])
        recent = tables.execute(sheets, "Sheet1", {
            "filters": [{"column": "Opened", "operator": "greater_or_equal", "value": "2024-01-01"},
                        {"column": "Balance", "operator": "less_than", "value": "1000"}],
            "metrics": [{"function": "count"}],
        })
        self.assertEqual(recent["rows"], [["2"]])
        listed = tables.execute(sheets, "Sheet1", {
            "filters": [{"column": "Branch", "operator": "one_of", "value": ["lae", "Madang"]}],
            "columns": ["Account No", "Balance"], "sortBy": "Balance", "sortDirection": "descending", "limit": 2,
        })
        self.assertEqual(listed["rows"], [["0012345", "1,200.50"], ["0012348", "500"]])
        self.assertEqual(listed["totalResultRows"], 4)

    def test_unsafe_plans_are_rejected(self) -> None:
        sheets = self.csv()
        for plan in (
            {"filters": [{"column": "Missing", "operator": "equals", "value": "x"}], "metrics": [{"function": "count"}]},
            {"metrics": [{"function": "sum", "column": "Branch"}]},
            {"filters": [{"column": "Branch", "operator": "greater_than", "value": "Lae"}], "metrics": [{"function": "count"}]},
            {"filters": [{"column": "Balance", "operator": "equals", "value": "lots"}], "metrics": [{"function": "count"}]},
        ):
            with self.assertRaises(tables.TablePlanError):
                tables.execute(sheets, "Sheet1", plan)

    def test_small_sheets_are_searchable_with_labelled_pages(self) -> None:
        text, labels, profiles = tables.sheet_text("loans.csv", self.csv())
        self.assertEqual(labels, ["Overview", "Rows 1–4"])
        self.assertIn("--- Page 2 ---", text)
        status = next(column for column in profiles[0]["columns"] if column["name"] == "Status")
        self.assertEqual(status["values"], [{"value": "Open", "count": 3}, {"value": "Closed", "count": 1}])


class OfficeDocumentTests(unittest.TestCase):
    def test_kind_uses_extension_for_inconsistent_browser_mime_types(self) -> None:
        self.assertEqual(office.document_kind("loans.csv", "application/vnd.ms-excel"), office.SPREADSHEET)
        self.assertEqual(office.document_kind("policy.docx", "application/octet-stream"), office.WORD)
        self.assertEqual(office.document_kind("scan.pdf", "application/pdf"), office.PDF)
        self.assertEqual(office.document_kind("photo.jpg", "image/jpeg"), office.IMAGE)
        self.assertIsNone(office.document_kind("deck.pptx", "application/octet-stream"))

    def test_word_headings_and_tables_keep_document_order(self) -> None:
        from docx import Document

        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "policy.docx"
            document = Document()
            document.add_heading("Eligibility", level=1)
            document.add_paragraph("Applicants must be salaried.")
            table = document.add_table(rows=2, cols=2)
            table.cell(0, 0).text, table.cell(0, 1).text = "Product", "Maximum"
            table.cell(1, 0).text, table.cell(1, 1).text = "Personal loan", "12 times gross monthly income"
            document.add_paragraph("Closing note.")
            document.save(path)
            text, page_count, labels = office.extract_word(path)
        self.assertEqual((page_count, labels), (1, ["Section 1"]))
        self.assertLess(text.index("# Eligibility"), text.index("Personal loan | 12 times"))
        self.assertLess(text.index("Personal loan"), text.index("Closing note."))

    def test_long_text_is_split_into_bounded_sections(self) -> None:
        paragraphs = [f"Paragraph {index}. " + "word " * 120 for index in range(12)]
        text, page_count, labels = office.sections_to_text(paragraphs)
        self.assertGreater(page_count, 1)
        self.assertEqual(labels[-1], f"Section {page_count}")
        self.assertIn("--- Page 2 ---", text)
        for section in text.split("--- Page ")[1:]:
            self.assertLessEqual(len(section), office.SECTION_CHARS + 20)


class SpreadsheetApiTests(unittest.TestCase):
    def setUp(self) -> None:
        self.temp = tempfile.TemporaryDirectory()
        root = Path(self.temp.name)
        self.patches = [
            patch.object(main, "DATABASE_PATH", root / "metadata.db"),
            patch.object(main, "UPLOADS_DIR", root / "uploads"),
            patch.object(main.rag, "index_document", lambda connection, upload_id, text, now: connection.execute(
                "INSERT OR REPLACE INTO document_indexes(upload_id, status, indexed_at) VALUES (?, 'ready', ?)", (upload_id, now))),
            patch.object(main.facts, "extract_document", side_effect=AssertionError("spreadsheets must not reach fact extraction")),
        ]
        for item in self.patches:
            item.start()
        main.initialise_storage()
        self.client = TestClient(main.app)

    def tearDown(self) -> None:
        for item in reversed(self.patches):
            item.stop()
        self.temp.cleanup()

    def upload_and_process(self, name: str = "loans.csv", content: bytes = LOAN_EXPORT.encode(), mime: str = "application/vnd.ms-excel", **form: str) -> str:
        response = self.client.post("/uploads", files={"file": (name, content, mime)}, data=form)
        self.assertEqual(response.status_code, 201)
        document = response.json()
        self.assertEqual(document["extractionStatus"], "queued")
        self.assertEqual(document["documentKind"], "spreadsheet")
        job = main.claim_next_job()
        main.process_job(job)
        return document["id"]

    def test_spreadsheet_is_profiled_without_fact_extraction(self) -> None:
        document_id = self.upload_and_process()
        extraction = self.client.get(f"/uploads/{document_id}/extraction").json()
        self.assertEqual(extraction["status"], "completed")
        self.assertEqual(extraction["pageLabels"], ["Overview", "Rows 1–4"])
        self.assertEqual(extraction["tables"][0]["rowCount"], 4)
        self.assertEqual(extraction["document"]["factStatus"], "not_applicable")
        retry = self.client.post(f"/uploads/{document_id}/facts/retry")
        self.assertEqual(retry.status_code, 400)

    def test_question_is_answered_exactly_from_rows(self) -> None:
        document_id = self.upload_and_process()
        plan = {"useTable": True, "tableId": f"{document_id}::Sheet1",
                "filters": [{"column": "Status", "operator": "equals", "value": "Open"}],
                "groupBy": ["Branch"], "metrics": [{"function": "count"}]}
        with patch.object(tables, "plan_query", return_value=plan) as planner:
            response = self.client.post("/questions", json={"question": "How many open loans per branch?", "documentIds": [document_id]})
        body = response.json()
        self.assertEqual(body["mode"], "table")
        self.assertEqual(body["table"], {"columns": ["Branch", "Count"], "rows": [["Madang", "2"], ["Lae", "1"]], "totalRows": 2})
        self.assertEqual(body["sources"][0]["pageLabel"], "3 of 4 rows")
        self.assertIn("Computed over 3 of 4 rows in loans.csv.", body["answer"])
        self.assertEqual(planner.call_args.args[2], "claude-haiku-4-5-20251001")

    def test_complete_plan_without_use_table_flag_still_runs(self) -> None:
        document_id = self.upload_and_process()
        plan = {"tableId": f"{document_id}::Sheet1", "metrics": [{"function": "sum", "column": "Balance"}]}
        with patch.object(tables, "plan_query", return_value=plan):
            body = self.client.post("/questions", json={"question": "Total balance?", "documentIds": [document_id]}).json()
        self.assertEqual(body["mode"], "table")
        self.assertEqual(body["table"]["rows"], [["2,000.50"]])

    def test_declined_or_invalid_plans_fall_back_to_document_search(self) -> None:
        document_id = self.upload_and_process()
        for plan in ({"useTable": False}, {"useTable": True, "tableId": f"{document_id}::Sheet1",
                                           "metrics": [{"function": "sum", "column": "Branch"}]}):
            with patch.object(tables, "plan_query", return_value=plan), \
                 patch.object(main.facts, "answer_with_trace", return_value=(None, {})), \
                 patch.object(main.rag, "retrieve_with_trace", return_value=([], {})):
                body = self.client.post("/questions", json={"question": "What does the policy say?", "documentIds": [document_id]}).json()
            self.assertEqual(body["mode"], "rag")
            self.assertIn("reason", body["debug"]["tables"])

    def test_agent_without_spreadsheet_capability_never_plans(self) -> None:
        document_id = self.upload_and_process()
        with patch.object(tables, "plan_query") as planner, \
             patch.object(main.rag, "retrieve_with_trace", return_value=([], {})):
            prepared, _chunks, trace = main._answer_from_scope("How many loans?", [document_id], {}, {"document_search"})
        planner.assert_not_called()
        self.assertEqual(trace["route"], "no-evidence")

    def test_chat_answer_keeps_its_result_table(self) -> None:
        document_id = self.upload_and_process(spaceId="portfolio-data")
        conversation = self.client.post("/conversations", json={"agentId": "portfolio-data-analyst"})
        self.assertEqual(conversation.status_code, 201, conversation.text)
        conversation_id = conversation.json()["id"]
        plan = {"useTable": True, "tableId": f"{document_id}::Sheet1", "metrics": [{"function": "count"}]}
        with patch.object(tables, "plan_query", return_value=plan), \
             patch.object(main.chat_history, "_claude_text", return_value="How many loans are there?"):
            stream = self.client.post(f"/conversations/{conversation_id}/messages/stream", json={"question": "How many loans are there?", "turnId": "t1"})
        self.assertIn("event: final", stream.text)
        answer = self.client.get(f"/conversations/{conversation_id}").json()["messages"][-1]
        self.assertEqual(answer["mode"], "table")
        self.assertEqual(answer["table"]["rows"], [["4"]])

    def test_previously_unsupported_office_files_are_queued_on_startup(self) -> None:
        with closing(main.get_connection()) as connection:
            connection.execute("INSERT INTO uploads(id, original_name, stored_name, mime_type, size_bytes, uploaded_at) VALUES ('old', 'notes.docx', 'old.docx', 'application/octet-stream', 1, '2026-01-01')")
            connection.execute("INSERT INTO extractions(upload_id, status, created_at, updated_at) VALUES ('old', 'unsupported', 'x', 'x')")
            connection.commit()
        main.initialise_storage()
        with closing(main.get_connection()) as connection:
            status = connection.execute("SELECT status FROM extractions WHERE upload_id='old'").fetchone()[0]
        self.assertEqual(status, "queued")

    def test_portfolio_agent_and_capability_are_available(self) -> None:
        catalog = self.client.get("/agents/catalog").json()
        self.assertIn("table_analysis", [item["id"] for item in catalog["capabilities"]])
        agent = self.client.get("/agents/portfolio-data-analyst").json()
        self.assertEqual(agent["spaceIds"], ["portfolio-data"])
        self.assertIn("table_analysis", agent["capabilities"])
        analyst = self.client.get("/agents/customer-document-analyst").json()
        self.assertIn("table_analysis", analyst["capabilities"])

    def test_deleting_a_spreadsheet_removes_its_profile(self) -> None:
        document_id = self.upload_and_process()
        self.assertEqual(self.client.delete(f"/uploads/{document_id}").status_code, 204)
        with closing(main.get_connection()) as connection:
            remaining = connection.execute("SELECT COUNT(*) FROM document_tables WHERE upload_id=?", (document_id,)).fetchone()[0]
        self.assertEqual(remaining, 0)


if __name__ == "__main__":
    unittest.main()
